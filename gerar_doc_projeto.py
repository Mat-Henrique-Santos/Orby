import os
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_docx(filename="Documentacao_Orby_Projeto_Integrado.docx"):
    doc = Document()

    # Margens ABNT
    for section in doc.sections:
        section.top_margin = Inches(1.18)
        section.left_margin = Inches(1.18)
        section.bottom_margin = Inches(0.79)
        section.right_margin = Inches(0.79)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # --- CAPA ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_inst.add_run("FACULDADE ANHANGUERA\nCURSO DE ANÁLISE E DESENVOLVIMENTO DE SISTEMAS\nDISCIPLINA: PROJETO INTEGRADO")
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph("\n" * 4)

    p_doc = doc.add_paragraph()
    p_doc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_doc = p_doc.add_run("DOCUMENTAÇÃO ESSENCIAL DO PROJETO DE SOFTWARE\n")
    run_doc.font.size = Pt(14)
    run_doc.font.bold = True
    run_doc.font.color.rgb = RGBColor(2, 132, 199)

    run_proj = p_doc.add_run("PROJETO: ORBY — GESTÃO OPERACIONAL DE ASSISTÊNCIA TÉCNICA")
    run_proj.font.size = Pt(12)
    run_proj.font.bold = True
    run_proj.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph("\n" * 4)

    p_integ = doc.add_paragraph()
    p_integ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_integ_title = p_integ.add_run("INTEGRANTES DO GRUPO\n")
    run_integ_title.font.size = Pt(11)
    run_integ_title.font.bold = True
    run_integ = p_integ.add_run("Matheus Henrique dos Santos\n(Líder do Projeto e Desenvolvedor)")
    run_integ.font.size = Pt(11)

    doc.add_paragraph("\n" * 4)

    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_footer = p_footer.add_run("POUSO ALEGRE - MG\n2026")
    run_footer.font.size = Pt(11)
    run_footer.font.bold = True

    doc.add_page_break()

    # --- IDENTIFICAÇÃO DO PROJETO ---
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(2, 132, 199)

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)

    add_heading_1("IDENTIFICAÇÃO DO PROJETO")

    table_id = doc.add_table(rows=10, cols=2)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_id.autofit = False

    data_id = [
        ("CAMPO", "PREENCHIMENTO"),
        ("Nome do projeto", "Orby — Gestão Operacional de Assistência Técnica"),
        ("Turma", "Análise e Desenvolvimento de Sistemas - ADS"),
        ("Semestre", "2º Semestre"),
        ("Professor", "Alex"),
        ("Líder ou responsável pelo contato", "Matheus Henrique dos Santos"),
        ("Integrantes e funções", "Matheus Henrique dos Santos — Líder de Projeto, Arquiteto de Software e Desenvolvedor"),
        ("Link do repositório", "https://github.com/mathe/orby"),
        ("Link do protótipo", "Executável Desktop Orby v2.0 (PySide6 / Windows x64)"),
        ("Data da versão final", "10 de Setembro de 2026"),
    ]

    for i, (campo, preench) in enumerate(data_id):
        row = table_id.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_margins(c0, 100, 100, 150, 150)
        set_cell_margins(c1, 100, 100, 150, 150)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(campo)
        r0.font.bold = True
        r0.font.size = Pt(10)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(preench)
        r1.font.size = Pt(10)

        if i == 0:
            set_cell_background(c0, "0284C7")
            set_cell_background(c1, "0284C7")
            r0.font.color.rgb = RGBColor(255, 255, 255)
            r1.font.color.rgb = RGBColor(255, 255, 255)
            r1.font.bold = True
        else:
            bg = "F1F5F9" if i % 2 == 1 else "FFFFFF"
            set_cell_background(c0, bg)
            set_cell_background(c1, bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # --- SUMÁRIO ---
    add_heading_1("SUMÁRIO")
    sumario_items = [
        ("IDENTIFICAÇÃO DO PROJETO", "2"),
        ("SUMÁRIO", "2"),
        ("1 DEFINIÇÃO DO PROJETO", "3"),
        ("   1.1 TEMA E CONTEXTO", "3"),
        ("   1.2 PROBLEMA", "3"),
        ("   1.3 SOLUÇÃO PROPOSTA", "3"),
        ("   1.4 OBJETIVO", "3"),
        ("2 USUÁRIOS E STAKEHOLDERS", "4"),
        ("   2.1 USUÁRIOS FINAIS", "4"),
        ("   2.2 STAKEHOLDERS", "4"),
        ("3 ESCOPO E MVP", "5"),
        ("   3.1 ESCOPO INCLUÍDO", "5"),
        ("   3.2 FORA DO ESCOPO", "5"),
        ("   3.3 PRODUTO MÍNIMO VIÁVEL", "5"),
        ("   3.4 RESTRIÇÕES", "6"),
        ("4 REQUISITOS DO SISTEMA", "6"),
        ("   4.1 REQUISITOS FUNCIONAIS", "6"),
        ("   4.2 REQUISITOS NÃO FUNCIONAIS", "7"),
        ("   4.3 REGRAS DE NEGÓCIO", "8"),
    ]

    p_sum_text = doc.add_paragraph()
    p_sum_text.paragraph_format.line_spacing = 1.2
    for item, pg in sumario_items:
        r_item = p_sum_text.add_run(f"{item.ljust(65, '.')} {pg}\n")
        r_item.font.size = Pt(10)
        if not item.startswith("   "):
            r_item.font.bold = True

    doc.add_page_break()

    # --- 1 DEFINIÇÃO DO PROJETO ---
    add_heading_1("1 DEFINIÇÃO DO PROJETO")
    
    add_heading_2("1.1 TEMA E CONTEXTO")
    p1 = doc.add_paragraph(
        "O presente projeto insere-se no setor de prestação de serviços de manutenção técnica especializada, "
        "com ênfase nas rotinas operacionais e financeiras de micro, pequenas e médias oficinas de assistência técnica de informática, "
        "dispositivos móveis e equipamentos eletroeletrônicos. Nesses estabelecimentos, a dinâmica de atendimento ao cliente "
        "no balcão envolve desde a triagem física inicial e o registro de acessórios até a elaboração de laudos diagnósticos, "
        "aprovação de orçamentos, execução de reparos, controle de peças, recebimento financeiro e devolução com garantia."
    )
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    p2 = doc.add_paragraph(
        "No contexto real investigado, constatou-se que grande parte das oficinas de pequeno porte opera de forma desestruturada, "
        "recorrendo a anotações em papel, conversas dispersas em aplicativos de mensagens ou planilhas eletrônicas desconectadas. "
        "Essa ausência de padronização compromete o controle do histórico de manutenções, dificulta a rastreabilidade dos números "
        "de série, gera frequentes atrasos na comunicação e ocasiona perdas financeiras por falta de cobrança correta de serviços executados e peças aplicadas."
    )
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_heading_2("1.2 PROBLEMA")
    p3 = doc.add_paragraph(
        "O problema central identificado é a vulnerabilidade operacional, a ausência de rastreabilidade no ciclo de vida das Ordens de Serviço (OS) "
        "e o descontrole no fluxo financeiro de caixa. Os principais afetados são os atendentes de recepção, que sofrem com cadastros duplicados e perda de dados; "
        "os técnicos de bancada, que perdem tempo buscando informações do defeito reclamado; e os proprietários, que não dispõem de visibilidade em tempo real sobre "
        "faturamento, ticket médio, margem de serviços versus peças e ordens em andamento."
    )
    p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p4 = doc.add_paragraph(
        "Essa desorganização acarreta extravio de componentes, duplicidade de cadastros, inconsistências contábeis e insatisfação dos clientes. "
        "Dessa forma, formula-se a seguinte questão norteadora: Como uma aplicação desktop verticalizada, ergonômica e dotada "
        "de validações transacionais rigorosas pode assegurar a integridade dos dados, otimizar o tempo de atendimento e garantir a rastreabilidade "
        "e o controle financeiro completo das ordens de serviço em pequenas assistências técnicas?"
    )
    p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_heading_2("1.3 SOLUÇÃO PROPOSTA")
    p5 = doc.add_paragraph(
        "A solução desenvolvida é o Orby, um sistema desktop de gestão operacional e financeira construído em Python 3.11+, PySide6 (Qt) e banco "
        "relacional SQLite local. O software atua como um 'mini-ERP verticalizado', oferecendo a robustez e a consistência de dados dos grandes "
        "sistemas corporativos, porém com interface visual limpa, intuitiva e sem os custos elevados ou a complexidade de implantação de plataformas web pesadas."
    )
    p5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p6 = doc.add_paragraph(
        "O sistema centraliza o cadastro de clientes com algoritmo inteligente de prevenção de duplicidades, vincula equipamentos com número de "
        "série e estado de entrada, e conduz cada Ordem de Serviço através de uma máquina de estados formal (Aberta, Em Diagnóstico, Aguardando Aprovação, "
        "Em Reparo, Aguardando Retirada, Finalizada e Cancelada). Além disso, dispõe de auditoria automática de status, importador de planilhas "
        "legadas (CSV/XLSX), módulo financeiro com discriminação de mão de obra e peças, e dashboard financeiro em tempo real com extrato de fechamento de caixa."
    )
    p6.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_heading_2("1.4 OBJETIVO")
    p7 = doc.add_paragraph(
        "Desenvolver e validar uma aplicação desktop de gerenciamento de ordens de serviço e controle financeiro para assistências técnicas, "
        "integrando cadastro com prevenção de duplicidades, rastreabilidade de equipamentos, esteira padronizada de status, cálculo de peças e serviços "
        "e importação de planilhas externas, a fim de eliminar inconsistências operacionais e proporcionar controle gerencial centralizado e seguro aos estabelecimentos."
    )
    p7.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # --- 2 USUÁRIOS E STAKEHOLDERS ---
    add_heading_1("2 USUÁRIOS E STAKEHOLDERS")
    
    add_heading_2("2.1 USUÁRIOS FINAIS")
    p_u = doc.add_paragraph(
        "O sistema é operado diretamente por diferentes perfis de colaboradores no ambiente da oficina, cada qual com permissões e "
        "necessidades específicas no fluxo de atendimento:"
    )
    p_u.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    t_users = doc.add_table(rows=4, cols=3)
    t_users.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_users.autofit = False

    u_data = [
        ("PERFIL DE USUÁRIO", "NECESSIDADE PRINCIPAL", "AÇÕES NO SISTEMA"),
        ("Atendente / Recepcionista", "Agilidade no atendimento de balcão, busca rápida de clientes, abertura de OS e lançamento de pagamentos.", "Cadastrar clientes e equipamentos, abrir Ordens de Serviço, consultar status por nome/telefone/documento, lançar recebimentos e importar planilhas legadas."),
        ("Técnico de Bancada", "Acesso claro ao defeito relatado, registro de diagnósticos, apontamento de peças e cálculo de mão de obra.", "Consultar fila de ordens, inserir laudos técnicos, atualizar status da OS (Em Diagnóstico/Reparo), discriminar valor de peças e serviços e registrar histórico."),
        ("Gerente / Proprietário", "Visão panorâmica do faturamento, controle de caixa por forma de pagamento e prevenção de perdas.", "Acompanhar métricas consolidadas no dashboard (faturamento realizado, previsão a receber, ticket médio), auditar histórico de status e autorizar exclusões."),
    ]

    for i, row_data in enumerate(u_data):
        row = t_users.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = [Inches(1.8), Inches(2.2), Inches(2.7)][j]
            set_cell_margins(cell, 90, 90, 120, 120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(cell, "0284C7")
            else:
                bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_heading_2("2.2 STAKEHOLDERS")
    p_s = doc.add_paragraph(
        "Além dos operadores diretos do software, o projeto impacta e recebe influência de diversas partes interessadas externas e organizacionais:"
    )
    p_s.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    t_stk = doc.add_table(rows=5, cols=3)
    t_stk.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_stk.autofit = False

    s_data = [
        ("STAKEHOLDER", "RELAÇÃO COM O PROJETO", "EXPECTATIVA OU INTERESSE"),
        ("Proprietário da Assistência", "Patrocinador e tomador de decisão da empresa.", "Aumentar a rentabilidade, reduzir o tempo médio de atendimento, evitar extravios de peças e ter controle de caixa sem custos recorrentes de nuvem."),
        ("Clientes da Assistência", "Consumidores finais dos serviços de reparo.", "Obter orçamentos transparentes, prazos cumpridos, histórico confiável de manutenções anteriores e segurança física de seus aparelhos."),
        ("Fornecedores de Peças", "Parceiros comerciais de suprimento de peças.", "Previsibilidade de demanda e clareza nas especificações técnicas das peças orçadas nas ordens de serviço."),
        ("Equipe de Dev (ADS)", "Responsáveis pelo ciclo de vida do software.", "Entregar um sistema robusto, modular, com código limpo, cobertura de testes automatizados e em conformidade com o Projeto Integrado."),
    ]

    for i, row_data in enumerate(s_data):
        row = t_stk.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = [Inches(1.8), Inches(2.2), Inches(2.7)][j]
            set_cell_margins(cell, 90, 90, 120, 120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(cell, "0284C7")
            else:
                bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)

    doc.add_page_break()

    # --- 3 ESCOPO E MVP ---
    add_heading_1("3 ESCOPO E MVP")

    add_heading_2("3.1 ESCOPO INCLUÍDO")
    p_e1 = doc.add_paragraph(
        "O escopo delimitado para o presente semestre contempla o núcleo transacional completo de atendimento e gestão operacional "
        "da assistência técnica, compreendendo os seguintes módulos e funcionalidades entregues:"
    )
    p_e1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    escopo_itens = [
        ("Módulo de Gestão de Clientes: ", "Cadastro completo com validação de dados obrigatórios, edição, exclusão segura com integridade referencial e checagem inteligente de duplicidades por documento ou nome/telefone."),
        ("Módulo de Gestão de Equipamentos: ", "Registro do parque de aparelhos por cliente, contendo tipo, marca, modelo, número de série/IMEI, avarias preexistentes, acessórios deixados e histórico de reparos."),
        ("Módulo de Ordens de Serviço (OS): ", "Abertura estruturada com vínculo obrigatório cliente-equipamento, máquina de estados com 7 etapas padronizadas, registro automático de histórico de transições e cálculo financeiro detalhado."),
        ("Módulo Financeiro e Fechamento de Caixa: ", "Painel financeiro com discriminação de mão de obra e peças, totalizadores de faturamento realizado, previsão a receber, ticket médio e distribuição por forma de pagamento (PIX, Cartão, Dinheiro)."),
        ("Dashboard de Indicadores Operacionais: ", "Painel analítico exibindo totalizadores de OS abertas, em reparo, finalizadas e faturamento acumulado, além de tabela das últimas ordens movimentadas."),
        ("Linha do Tempo Visual e Auditoria: ", "Visualizador cronológico dedicado na interface para auditoria detalhada de cada evento, alteração de status e laudo registrado na OS."),
        ("Módulo de Importação de Planilhas: ", "Mecanismo flexível para leitura de arquivos CSV e XLSX, mapeamento visual interativo de colunas, validação prévia de linhas e importação em lote."),
        ("Design System e Temas Visuais: ", "Interface desktop com estética moderna e ergonômica (Frutiger Aero / Glass), com suporte a alternância instantânea entre Modo Claro e Modo Escuro."),
    ]
    for tit, desc in escopo_itens:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(3)
        r1 = p_item.add_run(f"• {tit}")
        r1.font.bold = True
        r2 = p_item.add_run(desc)

    add_heading_2("3.2 FORA DO ESCOPO")
    p_fe = doc.add_paragraph(
        "Com o objetivo de assegurar a viabilidade técnica e a alta qualidade das entregas dentro do cronograma do semestre, os seguintes "
        "recursos foram formalmente declarados fora do escopo desta versão:"
    )
    p_fe.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    fora_itens = [
        ("Emissão Fiscal Automática: ", "Geração de notas fiscais de serviço eletrônicas (NFS-e, NF-e ou SAT) e integração com Secretarias da Fazenda."),
        ("Gateway de Pagamento Online: ", "Integração direta com adquirentes de cartão de crédito/débito via TEF ou liquidação automática de PIX via Webhook."),
        ("Sincronização Multi-Filiais em Nuvem: ", "Comunicação cliente-servidor distribuída via Internet para redes com múltiplas lojas físicas."),
        ("Notificação Automática por Mensageria: ", "Disparo automático de mensagens de aviso via SMS Gateway ou WhatsApp Business API."),
        ("Controle Contábil Complexo: ", "Módulos de folha de pagamento, depreciação patrimonial e escrituração contábil avançada."),
    ]
    for tit, desc in fora_itens:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(3)
        r1 = p_item.add_run(f"• {tit}")
        r1.font.bold = True
        r2 = p_item.add_run(desc)

    add_heading_2("3.3 PRODUTO MÍNIMO VIÁVEL (MVP)")
    p_mvp1 = doc.add_paragraph(
        "O Produto Mínimo Viável (MVP) do Orby constitui uma aplicação desktop autônoma e executável no Microsoft Windows, contendo o fluxo "
        "essencial e completo de operação de uma oficina: cadastro de cliente, associação de equipamento, abertura de OS com defeito relatado, "
        "tramitação de status de bancada, fechamento de valores com discriminação de peças e mão de obra, e fechamento de caixa com extrato financeiro."
    )
    p_mvp1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p_mvp2 = doc.add_paragraph(
        "Como diferencial do MVP, o sistema já incorpora travas de integridade relacional no SQLite (Foreign Keys ativas) para impedir a exclusão "
        "acidental de clientes ou aparelhos com histórico de atendimento, além da ferramenta de importação em lote de planilhas CSV/XLSX, permitindo "
        "que uma assistência técnica substitua controles legados de forma imediata e sem perda de dados históricos."
    )
    p_mvp2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_heading_2("3.4 RESTRIÇÕES")
    p_res = doc.add_paragraph(
        "O desenvolvimento e a implantação do Orby estão condicionados às seguintes restrições de projeto:"
    )
    p_res.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    restricoes = [
        ("Restrição de Plataforma e Execução: ", "A aplicação deve operar de modo 100% autônomo no sistema operacional Microsoft Windows (10 e 11 x64) e compatível nativamente com Linux, sem necessidade de instalação prévia de servidores externos."),
        ("Restrição de Stack Tecnológica: ", "Emprego exclusivo de Python 3.11+, PySide6 (Qt for Python) para renderização gráfica e SQLite 3 como motor de persistência em arquivo único local ('orby.db')."),
        ("Restrição Temporal e Acadêmica: ", "Conclusão das etapas de requisitos, modelagem, desenvolvimento e testes dentro do calendário letivo do 2º semestre da graduação em ADS na disciplina de Projeto Integrado."),
        ("Restrição de Desempenho em Balcão: ", "Inicialização em menos de 2 segundos e baixo consumo de memória RAM, garantindo fluidez mesmo em computadores básicos de recepção."),
    ]
    for tit, desc in restricoes:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(3)
        r1 = p_item.add_run(f"• {tit}")
        r1.font.bold = True
        r2 = p_item.add_run(desc)

    # --- 4 REQUISITOS DO SISTEMA ---
    add_heading_1("4 REQUISITOS DO SISTEMA")

    add_heading_2("4.1 REQUISITOS FUNCIONAIS")
    p_rf = doc.add_paragraph(
        "A Tabela a seguir detalha os requisitos funcionais que regem o comportamento e as capacidades operacionais e financeiras do sistema Orby, "
        "categorizados por código identificador, descrição da ação e nível de prioridade estabelecido pela equipe:"
    )
    p_rf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    t_rf = doc.add_table(rows=13, cols=3)
    t_rf.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_rf.autofit = False

    rf_data = [
        ("CÓDIGO", "REQUISITO FUNCIONAL", "PRIORIDADE"),
        ("RF01", "O sistema deverá permitir o cadastro, a consulta, a edição e a exclusão de clientes, armazenando nome completo, telefone, e-mail, CPF/CNPJ e endereço completo.", "Essencial"),
        ("RF02", "O sistema deverá validar a duplicidade de clientes no momento da inclusão ou importação, impedindo a inserção de registros com mesmo CPF/CNPJ ou mesmo telefone e nome.", "Essencial"),
        ("RF03", "O sistema deverá permitir o cadastro e a vinculação de equipamentos ao cliente, registrando tipo, marca, modelo, número de série/IMEI, avarias de entrada e acessórios deixados.", "Essencial"),
        ("RF04", "O sistema deverá permitir a exclusão de equipamentos, garantindo o bloqueio da operação caso existam Ordens de Serviço associadas ao item.", "Essencial"),
        ("RF05", "O sistema deverá permitir a abertura de Ordens de Serviço (OS), vinculando obrigatoriamente um cliente ativo, um equipamento cadastrado e a descrição do problema relatado.", "Essencial"),
        ("RF06", "O sistema deverá gerenciar os status da OS através de uma esteira controlada (Aberta, Em Diagnóstico, Aguardando Aprovação, Em Reparo, Aguardando Retirada, Finalizada e Cancelada).", "Essencial"),
        ("RF07", "O sistema deverá registrar automaticamente uma trilha de auditoria (historico_os) contendo a data/hora e o detalhamento de cada mudança de status ocorrida na OS.", "Essencial"),
        ("RF08", "O sistema deverá permitir a discriminação financeira do valor da OS em mão de obra (serviço) e peças aplicadas, calculando o valor total automaticamente.", "Essencial"),
        ("RF09", "O sistema deverá exigir a seleção da forma de pagamento (PIX, Dinheiro, Cartão de Crédito/Débito, Boleto) e valor maior que zero para a conclusão com status 'Finalizada'.", "Essencial"),
        ("RF10", "O sistema deverá disponibilizar um Módulo Financeiro com métricas consolidadas de faturamento realizado, previsão a receber, ticket médio e fechamento por forma de pagamento.", "Importante"),
        ("RF11", "O sistema deverá fornecer um assistente para importação em lote de clientes a partir de arquivos CSV e XLSX, com mapeamento visual de colunas e pré-visualização de validações.", "Importante"),
        ("RF12", "O sistema deverá disponibilizar um visualizador visual de Linha do Tempo da OS para acompanhamento cronológico de eventos e laudos técnicos.", "Importante"),
    ]

    for i, row_data in enumerate(rf_data):
        row = t_rf.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = [Inches(1.1), Inches(4.5), Inches(1.4)][j]
            set_cell_margins(cell, 80, 80, 110, 110)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(cell, "0284C7")
            else:
                bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                if j == 2:
                    r.font.bold = True
                    if text == "Essencial":
                        r.font.color.rgb = RGBColor(16, 185, 129)
                    elif text == "Importante":
                        r.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_page_break()

    # --- 4.2 REQUISITOS NÃO FUNCIONAIS ---
    add_heading_2("4.2 REQUISITOS NÃO FUNCIONAIS")
    p_rnf_desc = doc.add_paragraph(
        "Os requisitos não funcionais estabelecem os critérios de qualidade técnica, desempenho, confiabilidade e usabilidade "
        "que o Orby deve satisfazer para garantir uma operação eficiente e segura no balcão da assistência técnica:"
    )
    p_rnf_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    t_rnf = doc.add_table(rows=6, cols=4)
    t_rnf.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_rnf.autofit = False

    rnf_data = [
        ("CÓDIGO", "CATEGORIA", "REQUISITO VERIFICÁVEL", "PRIORIDADE"),
        ("RNF01", "Usabilidade", "O sistema deverá apresentar interface ergonômica com tema Frutiger Aero (Claro/Escuro), permitindo que um atendente realize cadastros e abertura de OS com tempo de aprendizado inferior a 15 minutos.", "Essencial"),
        ("RNF02", "Desempenho", "O sistema deverá responder a consultas e filtros de busca em menos de 500 milissegundos e inicializar a aplicação em menos de 2 segundos em computadores convencionais.", "Essencial"),
        ("RNF03", "Confiabilidade", "O sistema deverá assegurar atomicidade nas transações SQLite (commits/rollbacks) e integridade referencial com Foreign Keys ativas, impedindo a perda ou corrupção de dados em caso de queda de energia.", "Essencial"),
        ("RNF04", "Portabilidade", "O sistema deverá executar de forma autônoma no Microsoft Windows (10/11 x64) via executável compilado e ser compatível nativamente com ambientes Linux (X11/Wayland).", "Essencial"),
        ("RNF05", "Segurança Local", "O sistema deverá armazenar a base de dados em arquivo local único ('orby.db'), sem tráfego de dados de clientes ou ordens de serviço para servidores externos não autorizados.", "Importante"),
    ]

    for i, row_data in enumerate(rnf_data):
        row = t_rnf.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = [Inches(1.0), Inches(1.4), Inches(3.4), Inches(1.2)][j]
            set_cell_margins(cell, 80, 80, 110, 110)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(cell, "0284C7")
            else:
                bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                if j == 3:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(16, 185, 129) if text == "Essencial" else RGBColor(2, 132, 199)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- 4.3 REGRAS DE NEGÓCIO ---
    add_heading_2("4.3 REGRAS DE NEGÓCIO")
    p_rn_desc = doc.add_paragraph(
        "As regras de negócio definem as restrições, validações e comportamentos que governam os processos operacionais "
        "e garantem a integridade dos dados no Orby:"
    )
    p_rn_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    t_rn = doc.add_table(rows=9, cols=3)
    t_rn.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_rn.autofit = False

    rn_data = [
        ("CÓDIGO", "REGRA DE NEGÓCIO", "REQUISITO RELACIONADO"),
        ("RN01", "Todo equipamento cadastrado deve pertencer compulsoriamente a um cliente ativo cadastrado no sistema.", "RF03, RF04"),
        ("RN02", "Uma Ordem de Serviço não pode ser criada sem a indicação de um cliente ativo, um equipamento vinculado e a descrição do problema.", "RF05"),
        ("RN03", "A transição de uma OS para o status 'Finalizada' exige obrigatoriamente valor total maior que zero e a seleção de uma forma de pagamento válida.", "RF08, RF09"),
        ("RN04", "É proibida a exclusão de clientes ou equipamentos que possuam vínculos com Ordens de Serviço ativas ou históricas.", "RF01, RF04"),
        ("RN05", "Toda alteração de status na esteira da OS deve gerar um registro imutável na tabela de auditoria (historico_os) com timestamp e status anterior/novo.", "RF06, RF07"),
        ("RN06", "O sistema deve impedir a inserção de clientes em duplicidade, considerando mesmo CPF/CNPJ ou combinação idêntica de nome completo e telefone.", "RF02, RF11"),
        ("RN07", "O valor total da Ordem de Serviço deve ser calculado pela soma do valor da mão de obra (serviço) com o valor das peças aplicadas.", "RF08, RF10"),
        ("RN08", "A exclusão física de uma Ordem de Serviço é uma operação restrita que remove em cascata o histórico de auditoria associado.", "RF05, RF07"),
    ]

    for i, row_data in enumerate(rn_data):
        row = t_rn.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = [Inches(1.1), Inches(4.5), Inches(1.4)][j]
            set_cell_margins(cell, 80, 80, 110, 110)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(cell, "0284C7")
            else:
                bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                if j == 2:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(2, 132, 199)

    doc.save(filename)
    print(f"DOCX atualizado com sucesso: {filename}")

if __name__ == "__main__":
    create_docx()
