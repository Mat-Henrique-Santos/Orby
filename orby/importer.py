from __future__ import annotations

import csv
from pathlib import Path

def read_table(path: str) -> tuple[list[str], list[dict[str, str]]]:
    file_path = Path(path)
    if file_path.suffix.lower() == ".csv":
        with file_path.open("r", encoding="utf-8-sig", newline="") as file:
            sample = file.read(4096)
            file.seek(0)
            delimiter = ","
            try:
                sniffer = csv.Sniffer()
                dialect = sniffer.sniff(sample, delimiters=";,|\t,")
                delimiter = dialect.delimiter
            except Exception:
                if ";" in sample and "," not in sample:
                    delimiter = ";"
            reader = csv.DictReader(file, delimiter=delimiter)
            headers = [h.strip() for h in (reader.fieldnames or []) if h]
            rows = [{key.strip(): (value or "").strip() for key, value in row.items() if key} for row in reader]
            return headers, rows
    from openpyxl import load_workbook
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    rows = []
    for data in values[1:]:
        if any(value not in (None, "") for value in data):
            rows.append({headers[index]: str(value).strip() if value is not None else "" for index, value in enumerate(data)})
    return headers, rows


from .core.models import validar_documento, validar_email, validar_telefone


def validate_customers(rows: list[dict[str, str]], mapping: dict[str, str], duplicate_check) -> tuple[list[dict], list[dict]]:
    approved, rejected, seen = [], [], set()
    for index, source in enumerate(rows, start=2):
        item = {field: source.get(column, "").strip() for field, column in mapping.items() if column}
        errors = []
        if not item.get("nome"):
            errors.append("Nome obrigatório")
        if not item.get("telefone"):
            errors.append("Telefone obrigatório")
        elif not validar_telefone(item["telefone"]):
            errors.append("Telefone inválido (10 ou 11 dígitos)")

        if item.get("email") and not validar_email(item["email"]):
            errors.append("E-mail com formato inválido")

        if item.get("documento") and not validar_documento(item["documento"]):
            errors.append("Documento inválido (CPF/CNPJ)")

        key = (item.get("nome", "").casefold(), item.get("telefone", ""))
        is_dup = False
        if key in seen:
            is_dup = True
        elif item.get("nome") and item.get("telefone") and duplicate_check:
            try:
                is_dup = duplicate_check(item["nome"], item["telefone"], item.get("documento", ""))
            except TypeError:
                is_dup = duplicate_check(item["nome"], item["telefone"])

        if is_dup:
            errors.append("Possível duplicidade")
        seen.add(key)
        if errors:
            rejected.append({"linha": index, "dados": item, "motivo": "; ".join(errors)})
        else:
            approved.append(item)
    return approved, rejected
